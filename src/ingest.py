# ingest.py
from pathlib import Path
import fitz  # PyMuPDF
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def extrair_pdf(caminho: Path):
    doc = fitz.open(caminho)
    itens = []
    for n_pag, page in enumerate(doc):
        # --- imagens embutidas ---
        for img in page.get_images(full=True):
            xref = img[0]
            rects = page.get_image_rects(xref)
            bbox = rects[0] if rects else None
            pix = fitz.Pixmap(doc, xref)
            if pix.n > 4:  # CMYK → RGB
                pix = fitz.Pixmap(fitz.csRGB, pix)
            itens.append({
                "tipo": "imagem",
                "pagina": n_pag,
                "bbox": bbox,
                "bytes": pix.tobytes("png"),
            })
        # --- texto com posição (legendas) ---
        for bloco in page.get_text("dict")["blocks"]:
            if bloco.get("type") == 0:  # texto
                texto = " ".join(
                    span["text"]
                    for linha in bloco["lines"]
                    for span in linha["spans"]
                ).strip()
                if texto:
                    itens.append({
                        "tipo": "texto",
                        "pagina": n_pag,
                        "bbox": fitz.Rect(bloco["bbox"]),
                        "texto": texto,
                    })
        # --- página renderizada (fallback p/ recorte) ---
        itens.append({
            "tipo": "pagina_render",
            "pagina": n_pag,
            "pix": page.get_pixmap(dpi=200),
        })
    return itens


def extrair_pptx(caminho: Path):
    prs = Presentation(caminho)
    itens = []
    for n_slide, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                itens.append({
                    "tipo": "imagem",
                    "pagina": n_slide,
                    "bbox": (shape.left, shape.top,
                             shape.left + shape.width,
                             shape.top + shape.height),
                    "bytes": shape.image.blob,
                })
            elif shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    itens.append({
                        "tipo": "texto",
                        "pagina": n_slide,
                        "bbox": (shape.left, shape.top,
                                 shape.left + shape.width,
                                 shape.top + shape.height),
                        "texto": txt,
                    })
    return itens